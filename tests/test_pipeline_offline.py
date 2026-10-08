"""Offline regression tests (no API key, no model download).

Run:  pytest -q
Uses data/input/sample.pdf (the 11-page archive).
"""
from pathlib import Path

import pytest

from ingestion import load_document
from preprocessing.cleaner import clean_documents
from preprocessing.chunker import chunk_documents
from rag.embeddings import get_max_tokens, get_token_counter
from rag.retriever import get_retriever
from rag.vectorstore import create_vectorstore
from extraction.table_parser import parse_date, parse_tables
from extraction.extractor import create_extractor
from output.txt_writer import write_txt

SAMPLE = Path(__file__).resolve().parents[1] / "data" / "input" / "sample.pdf"
pytestmark = pytest.mark.skipif(not SAMPLE.exists(), reason="sample.pdf missing")


@pytest.fixture(scope="module")
def docs():
    return clean_documents(load_document(str(SAMPLE)))


@pytest.fixture(scope="module")
def chunks(docs):
    return chunk_documents(docs)


@pytest.fixture(scope="module")
def retriever(chunks):
    return get_retriever(create_vectorstore(chunks), k=3)


def test_split_tables_are_rejoined(docs):
    by_title = {d.metadata["title"]: d.metadata for d in docs if d.metadata["element_type"] == "table"}
    master = by_title["1. MASTER RECORD EXTRACT"]
    assert master["n_rows"] == 64 and master["pages"] == [1, 2]
    monthly = by_title["2. MONTHLY PERFORMANCE TABLE — EXTRACT A"]
    assert monthly["n_rows"] == 12 and monthly["pages"] == [3, 4]      # Jan-Jul + Aug-Dec
    kv = by_title["5. RAW KEY/VALUE EXPORT"]
    assert kv["n_rows"] == 55 and not kv["has_header"]                 # header-less, 2 pages


def test_side_tables_not_merged_across_pages(docs):
    titles = [d.metadata["title"] for d in docs if d.metadata["element_type"] == "table"]
    assert "RELEASE CHECKPOINT | P03 TABLE 03" in titles
    assert "RELEASE CHECKPOINT | P04 TABLE 01" in titles


def test_reading_order_and_footer_removal(docs):
    text = "\n".join(d.page_content for d in docs)
    assert "Synthetic Test Dataset" not in text           # footer stripped
    assert "Aug 18 | Delivery | Site code GGN" in text


def test_chunks_never_exceed_embedding_window(chunks):
    count = get_token_counter()
    assert max(count(c.page_content) for c in chunks) <= get_max_tokens()


def test_table_chunks_cut_on_row_boundaries_and_repeat_header(chunks):
    master = [c for c in chunks if c.metadata.get("title") == "1. MASTER RECORD EXTRACT"]
    assert len(master) > 1
    for c in master:
        assert "Columns: ID | Date" in c.page_content
        for line in c.page_content.splitlines():
            if line.startswith("R-"):
                assert line.count(" | ") == 8                 # complete row, 9 cells
    covered = sum(c.metadata["row_end"] - c.metadata["row_start"] + 1 for c in master)
    assert covered == 64


@pytest.mark.parametrize("query,needle", [
    ("Which site code appears as Gurgaon, Gurugram and GGM?", "GGN appears as Gurgaon"),
    ("invoice batch 77B has 214 records but reconciliation says 219", "214 records"),
    ("R-1841-A duplicate child record", "R-1841-A"),
    ("R-1058", "R-1058 |"),
    ("Oct data gap CSAT variance", "Oct | 2008"),
    ("DQ-007 visual-only value", "DQ-007"),
    ("bundle-09-3 checksum", "32FDFA"),
])
def test_targeted_retrieval_finds_the_exact_chunk(retriever, query, needle):
    assert any(needle in d.page_content for d in retriever.invoke(query))


def test_table_parser_counts(docs):
    parsed = parse_tables(docs)
    assert len(parsed["master_records"]) == 64
    assert len(parsed["monthly_performance"]) == 12
    assert len(parsed["data_quality_flags"]) == 8
    assert {g.title for g in parsed["value_groups"]} >= {"Traffic distribution"}
    ids = [r.record_id for r in parsed["master_records"]]
    assert ids == [f"R-{n}" for n in range(1001, 1065)]
    first = parsed["master_records"][0]
    assert (first.metric_a, first.metric_b, first.owner) == (4524, 24.6, None)


def test_ancillary_tables_keep_all_rows_and_render_in_txt(docs, tmp_path):
    parsed = parse_tables(docs)
    ancillary = parsed["ancillary_tables"]
    by_name = {table.table_name: table for table in ancillary}

    archive = by_name["ARCHIVE MANIFEST | P03 TABLE 06"]
    assert archive.page_number == 3
    assert archive.headers == ["Bundle name", "Created date", "Checksum fragment"]
    assert len(archive.rows) == 4
    assert archive.rows[0] == ["bundle-03-1", "2026-10-16", "17711E"]

    side_titles = {
        doc.metadata["title"] for doc in docs
        if doc.metadata.get("element_type") == "table" and doc.metadata.get("kind") == "side"
    }
    assert side_titles <= set(by_name)

    extracted = create_extractor()(docs, {}, source_file=str(SAMPLE))
    assert len(extracted.ancillary_tables) == len(ancillary)
    assert extracted.report["ancillary_tables"] == len(ancillary)

    output = tmp_path / "extracted.txt"
    write_txt(extracted, output)
    text = output.read_text(encoding="utf-8")
    assert "Other Tables:" not in text
    for table in extracted.ancillary_tables:
        assert f"{table.table_name} (Page {table.page_number})" in text
        for row in table.rows:
            assert " | ".join(row) in text


def test_dates_are_only_normalised_when_unambiguous():
    assert parse_date("2026-05-17") == ("2026-05-17", None)
    assert parse_date("17 May 2026") == ("2026-05-17", None)
    assert parse_date("06-21-2026") == ("2026-06-21", None)
    assert parse_date("03/04/26") == (None, "date_ambiguous")
    assert parse_date("Aug 3") == (None, "date_missing_year")


def test_xlsx_and_docx_loaders(tmp_path):
    import openpyxl
    from docx import Document
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ops"
    ws.append(["Queue", "Group", "Hours"]); ws.append(["Q1", "Billing", 4]); ws.append(["Q2", "SRE", 7.5])
    wb.save(tmp_path / "t.xlsx")
    xl = load_document(str(tmp_path / "t.xlsx"))
    assert xl[0].metadata["rows"] == [["Q1", "Billing", "4"], ["Q2", "SRE", "7.5"]]
    assert "Q1 | Billing | 4" in xl[0].page_content            # pipe-separated, not padded

    d = Document(); d.add_heading("Results", 1); d.add_paragraph("Table 1: Sales")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Region", "Sales", "North", "10"
    d.save(tmp_path / "t.docx")
    dx = [x for x in load_document(str(tmp_path / "t.docx")) if x.metadata["element_type"] == "table"]
    assert dx[0].metadata["title"] == "Table 1: Sales" and dx[0].metadata["rows"] == [["North", "10"]]

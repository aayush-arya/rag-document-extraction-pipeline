"""Lossless, layout-aware loader for Excel workbooks."""

from __future__ import annotations

import datetime as dt
import posixpath
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
from langchain_core.documents import Document
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from .table_utils import RawTable, make_unique_columns, norm_name, table_to_document


_DATE_HEADERS = {
    "date", "event date", "created date", "updated date", "last updated",
    "timestamp", "created at", "updated at", "datetime", "date time",
}


def _xml_numeric_values(file_path: str):
    """Retain the exact numeric lexical values stored in worksheet XML."""
    values_by_sheet = {}
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(file_path) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relations_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in relations_root.findall("pkg:Relationship", ns)
        }
        for sheet in workbook_root.findall("main:sheets/main:sheet", ns):
            target = targets.get(sheet.attrib.get(f"{{{ns['rel']}}}id"))
            if not target:
                continue
            sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
                posixpath.join("xl", target)
            )
            sheet_root = ElementTree.fromstring(archive.read(sheet_path))
            cells = {}
            for cell in sheet_root.findall(".//main:c", ns):
                if cell.attrib.get("t", "n") != "n":
                    continue
                raw = cell.findtext("main:v", namespaces=ns)
                if raw is not None:
                    cells[cell.attrib["r"]] = raw
            values_by_sheet[sheet.attrib["name"]] = cells
    return values_by_sheet


def _format_value(value, cell, date_hint: bool, epoch) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        if isinstance(value, dt.datetime):
            if date_hint:
                return value.date().isoformat()
            return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, Decimal):
        if cell.is_date or date_hint:
            try:
                converted = from_excel(float(value), epoch)
            except (OverflowError, ValueError):
                converted = value
            if isinstance(converted, (dt.datetime, dt.date)):
                return converted.date().isoformat() if isinstance(converted, dt.datetime) else converted.isoformat()
        number = format(value, "f")
        if "%" in (cell.number_format or ""):
            number = format(value * 100, "f")
            if "." in number:
                number = number.rstrip("0").rstrip(".")
            return (number or "0") + "%"
        return number
    if isinstance(value, (int, float)):
        if cell.is_date or date_hint:
            try:
                converted = from_excel(value, epoch)
            except (OverflowError, ValueError):
                converted = value
            if isinstance(converted, (dt.datetime, dt.date)):
                return converted.date().isoformat() if isinstance(converted, dt.datetime) else converted.isoformat()
        number = str(value)
        if "%" in (cell.number_format or ""):
            number = format(Decimal(number) * 100, "f")
            if "." in number:
                number = number.rstrip("0").rstrip(".")
            return number + "%"
        return number
    return str(value).strip()


def _load_value_sheet(file_path: str):
    keep_vba = Path(file_path).suffix.lower() == ".xlsm"
    values = openpyxl.load_workbook(file_path, data_only=True, keep_vba=keep_vba)
    formulas = openpyxl.load_workbook(file_path, data_only=False, keep_vba=keep_vba)
    return values, formulas


def _worksheet_grid(value_sheet, formula_sheet, epoch, numeric_xml_values):
    """Build a sparse grid, expanding merged ranges and restoring absent formula caches."""
    coordinates = {}
    actual_coordinates = set()
    max_row = max(value_sheet.max_row, formula_sheet.max_row)
    max_column = max(value_sheet.max_column, formula_sheet.max_column)

    for row in value_sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
        for cell in row:
            value = cell.value
            formula_value = formula_sheet[cell.coordinate].value
            if value is None and isinstance(formula_value, str) and formula_value.startswith("="):
                value = formula_value
            elif cell.data_type == "n" and cell.coordinate in numeric_xml_values:
                try:
                    value = Decimal(numeric_xml_values[cell.coordinate])
                except InvalidOperation:
                    pass
            if value is not None:
                coordinates[(cell.row, cell.column)] = (value, cell, True)
                actual_coordinates.add((cell.row, cell.column))

    for merged_range in value_sheet.merged_cells.ranges:
        anchor = value_sheet.cell(merged_range.min_row, merged_range.min_col)
        value = anchor.value
        if value is None:
            formula_value = formula_sheet[anchor.coordinate].value
            value = formula_value if isinstance(formula_value, str) and formula_value.startswith("=") else None
        if value is None:
            continue
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                coordinates[(row, column)] = (value, anchor, (row, column) == (merged_range.min_row, merged_range.min_col))
        actual_coordinates.add((merged_range.min_row, merged_range.min_col))
        max_row = max(max_row, merged_range.max_row)
        max_column = max(max_column, merged_range.max_col)

    if not coordinates:
        return {}, set()
    return coordinates, actual_coordinates


def _row_blocks(coordinates):
    occupied_rows = sorted({row for row, _ in coordinates})
    blocks, rows, previous = [], [], None
    for row in occupied_rows:
        if previous is not None and row > previous + 1:
            blocks.append(rows)
            rows = []
        rows.append(row)
        previous = row
    if rows:
        blocks.append(rows)
    return blocks


def _column_blocks(rows, coordinates):
    """Split a populated row band at columns empty throughout the whole band."""
    row_set = set(rows)
    occupied_columns = sorted({column for row, column in coordinates if row in row_set})
    groups, columns, previous = [], [], None
    for column in occupied_columns:
        if previous is not None and column > previous + 1:
            groups.append(columns)
            columns = []
        columns.append(column)
        previous = column
    if columns:
        groups.append(columns)
    return groups


def _value_at(coordinates, row, column, date_hint, epoch):
    item = coordinates.get((row, column))
    if item is None:
        return ""
    value, cell, _ = item
    return _format_value(value, cell, date_hint, epoch)


def _header_row_index(rows, columns, coordinates):
    """Find a sparse header row near the top without treating numeric data as labels."""
    for row in rows:
        values = [
            str(coordinates[(row, column)][0]).strip()
            for column in columns if (row, column) in coordinates
        ]
        if not values:
            continue
        if any(isinstance(coordinates[(row, column)][0], (int, float, Decimal, dt.date, dt.datetime))
               for column in columns if (row, column) in coordinates):
            return None
        names = {norm_name(value) for value in values}
        if len(values) >= 2 or names.intersection(_DATE_HEADERS | {
            "id", "record id", "status", "owner", "month", "period", "requests",
            "resolved", "amount", "name", "department", "location",
        }):
            return row
    return None


def _merged_title(rows, columns, coordinates, worksheet):
    if len(rows) < 2:
        return None, rows
    first = rows[0]
    merged_titles = [
        merged for merged in worksheet.merged_cells.ranges
        if merged.min_row == first and merged.max_row == first
        and merged.min_col <= min(columns) and merged.max_col >= max(columns)
    ]
    if len(merged_titles) == 1:
        value = coordinates.get((first, merged_titles[0].min_col), (None, None, None))[0]
        return (str(value).strip() if value is not None else None), rows[1:]
    return None, rows


def _merged_headers(header_rows, columns, coordinates, epoch, date_hints):
    headers = []
    for column in columns:
        parts = []
        for row in header_rows:
            value = _value_at(coordinates, row, column, False, epoch).strip()
            if value and (not parts or parts[-1] != value):
                parts.append(value)
        headers.append(" / ".join(parts))
    return make_unique_columns(headers)


def _block_metadata(sheet, sheet_index, file_path, rows, columns):
    sheet_name = sheet.title.replace("'", "''")
    quoted_name = f"'{sheet_name}'" if re.search(r"[^A-Za-z0-9_]", sheet.title) else sheet_name
    block_range = (
        f"{quoted_name}!{get_column_letter(min(columns))}{min(rows)}:"
        f"{get_column_letter(max(columns))}{max(rows)}"
    )
    return {
        "source": file_path,
        "file_name": Path(file_path).name,
        "file_type": Path(file_path).suffix.lower().lstrip("."),
        "sheet": sheet.title,
        "sheet_name": sheet.title,
        "sheet_index": sheet_index,
        "sheet_state": sheet.sheet_state,
        "block_range": block_range,
        "section": sheet.title,
    }


def load_excel(file_path: str) -> list[Document]:
    """Read every worksheet and emit independent text blocks / structured tables."""
    value_workbook, formula_workbook = _load_value_sheet(file_path)
    documents: list[Document] = []
    table_count = 0
    extension = Path(file_path).suffix.lower().lstrip(".")
    numeric_values = _xml_numeric_values(file_path)

    for sheet_index, (sheet, formula_sheet) in enumerate(
        zip(value_workbook.worksheets, formula_workbook.worksheets), start=1
    ):
        coordinates, actual_coordinates = _worksheet_grid(
            sheet, formula_sheet, value_workbook.epoch, numeric_values.get(sheet.title, {})
        )
        if not coordinates:
            continue

        pending_title = None
        pending_title_row = None
        pending_title_columns = []
        for row_group in _row_blocks(coordinates):
            for column_group in _column_blocks(row_group, coordinates):
                base_meta = _block_metadata(sheet, sheet_index, file_path, row_group, column_group)
                row_set, column_set = set(row_group), set(column_group)
                populated_columns = {
                    column for row, column in actual_coordinates
                    if row in row_set and column in column_set
                }

                if len(row_group) == 1 and len(column_group) > 1:
                    row = row_group[0]
                    values_in_row = {
                        str(coordinates[(row, column)][0]).strip()
                        for column in column_group if (row, column) in coordinates
                    }
                    is_merged_title = any(
                        merged.min_row == row and merged.max_row == row
                        and merged.max_col > merged.min_col
                        and merged.min_col <= min(column_group)
                        and merged.max_col >= max(column_group)
                        for merged in sheet.merged_cells.ranges
                    )
                    if is_merged_title and len(values_in_row) == 1:
                        pending_title = next(iter(values_in_row))
                        pending_title_row = row
                        pending_title_columns = column_group
                        continue

                if len(populated_columns) <= 1:
                    lines = []
                    first_value = _value_at(
                        coordinates, row_group[0], column_group[0], False, value_workbook.epoch
                    )
                    is_date_column = norm_name(first_value) in _DATE_HEADERS
                    for row in row_group:
                        values = [
                            _value_at(
                                coordinates, row, column,
                                is_date_column and row != row_group[0],
                                value_workbook.epoch,
                            ) for column in column_group
                        ]
                        unique = list(dict.fromkeys(value for value in values if value))
                        if unique:
                            lines.append(" ".join(unique))
                    text = "\n".join(lines).strip()
                    if text:
                        documents.append(Document(page_content=text, metadata={
                            **base_meta, "element_type": "text", "is_heading": False,
                            "region": "body", "page": sheet_index, "pages": [sheet_index],
                        }))
                    continue

                title, content_rows = _merged_title(row_group, column_group, coordinates, sheet)
                if (pending_title and pending_title_row is not None
                        and min(row_group) - pending_title_row <= 3
                        and set(pending_title_columns).intersection(column_group)):
                    title = title or pending_title
                    pending_title = None
                    pending_title_row = None
                    pending_title_columns = []
                header_row = _header_row_index(content_rows, column_group, coordinates)
                header_rows = [header_row] if header_row is not None else []
                if header_row is not None:
                    has_group_header = any(
                        merged.min_row == header_row and merged.max_row == header_row
                        and merged.max_col > merged.min_col
                        and merged.min_col <= max(column_group)
                        and merged.max_col >= min(column_group)
                        for merged in sheet.merged_cells.ranges
                    )
                    for candidate in content_rows:
                        if candidate != header_row + 1 or not has_group_header:
                            continue
                        candidate_values = [
                            coordinates[(candidate, column)][0]
                            for column in column_group if (candidate, column) in coordinates
                        ]
                        if candidate_values and not any(
                            isinstance(value, (int, float, Decimal, dt.date, dt.datetime))
                            for value in candidate_values
                        ):
                            header_rows.append(candidate)
                    header_rows.sort()

                headers = _merged_headers(
                    header_rows, column_group, coordinates, value_workbook.epoch, [False] * len(column_group)
                )
                data_rows = [row for row in content_rows if row not in header_rows]
                date_hints = [
                    any(norm_name(label) in _DATE_HEADERS for label in header.split(" / "))
                    for header in headers
                ]
                values = [
                    [_value_at(coordinates, row, column, date_hints[index], value_workbook.epoch)
                     for index, column in enumerate(column_group)]
                    for row in data_rows
                ]

                if header_row is None:
                    headers = [f"col_{index + 1}" for index in range(len(column_group))]
                else:
                    headers = make_unique_columns(headers)
                nonempty_rows = [row for row in values if any(value for value in row)]
                if not nonempty_rows:
                    continue

                table_count += 1
                table = RawTable(
                    rows=nonempty_rows,
                    columns=headers,
                    has_header=header_row is not None,
                    title=title or base_meta["block_range"],
                    section=sheet.title,
                    pages=[sheet_index],
                    sheet=sheet.title,
                    header_row=headers if header_row is not None else None,
                    paged=False,
                    file_name=Path(file_path).name,
                    block_range=base_meta["block_range"],
                    sheet_index=sheet_index,
                )
                document = table_to_document(table, file_path, extension, f"t{table_count:03d}")
                document.metadata.update(base_meta)
                document.metadata.update({
                    "page": sheet_index, "pages": [sheet_index],
                    "headers": headers, "rows": nonempty_rows,
                    "secondary_sheet": sheet_index > 1,
                })
                documents.append(document)

    for workbook in (value_workbook, formula_workbook):
        workbook.close()
        vba_archive = getattr(workbook, "vba_archive", None)
        if vba_archive is not None:
            vba_archive.close()
    return documents

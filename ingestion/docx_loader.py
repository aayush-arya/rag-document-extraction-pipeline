"""Word (.docx) loader.

* walks the body in true document order (paragraphs and tables interleaved)
* headings (Heading n / Title / short all-bold lines) become the running
  ``section``; the short paragraph right before a table is used as its title
* tables stay structured (merged cells de-duplicated, header row from the
  ``tblHeader`` flag or the usual heuristic) and a table that continues
  directly into the next one with the same header is stitched back together
* embedded pictures are sent through the OCR / vision fallback
"""

from __future__ import annotations

import io
import logging

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document

from .ocr import ocr_image
from .table_utils import (
    RawTable, clean_cell, make_unique_columns, norm_name, split_table, table_to_document,
)

logger = logging.getLogger(__name__)

TITLE_MAX_CHARS = 100


def _is_heading(paragraph: Paragraph) -> bool:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style.startswith("Heading") or style == "Title":
        return True
    runs = [r for r in paragraph.runs if r.text.strip()]
    text = paragraph.text.strip()
    return bool(runs) and all(r.bold for r in runs) and len(text) <= TITLE_MAX_CHARS


def _is_caption(paragraph: Paragraph) -> bool:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    return style == "Caption" or paragraph.text.strip().lower().startswith("table")


def _table_rows(table: Table) -> tuple[list[list[str]], bool]:
    """Cell grid with merged cells de-duplicated + whether row 1 is flagged as header."""
    rows = []
    for row in table.rows:
        cells, last_tc = [], None
        for cell in row.cells:
            if cell._tc is last_tc:        # horizontally merged cell appears repeatedly
                continue
            last_tc = cell._tc
            cells.append(clean_cell("\n".join(p.text for p in cell.paragraphs)))
        rows.append(cells)
    flagged = False
    if table.rows:
        trpr = table.rows[0]._tr.trPr
        flagged = trpr is not None and any(c.tag.endswith("}tblHeader") for c in trpr)
    return rows, flagged


def _text_doc(text, source, section, heading=False, element_type="text", **extra) -> Document:
    metadata = {
        "source": source, "file_type": "docx", "element_type": element_type,
        "page": 1, "pages": [1], "section": section, "is_heading": heading, "region": "body",
        "ocr": element_type == "ocr_text",
    }
    metadata.update(extra)
    return Document(page_content=text, metadata=metadata)


def load_docx(file_path: str) -> list[Document]:
    doc = DocxDocument(file_path)
    elements: list = []
    buffer: list[Paragraph] = []
    section = None

    def flush():
        nonlocal buffer
        text = "\n".join(p.text.strip() for p in buffer if p.text.strip())
        if text:
            elements.append(_text_doc(text, file_path, section))
        buffer = []

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            if _is_heading(paragraph):
                flush()
                section = text
                elements.append(_text_doc(text, file_path, section, heading=True))
            else:
                buffer.append(paragraph)

        elif tag == "tbl":
            raw_rows, flagged = _table_rows(Table(child, doc))
            parsed = split_table(raw_rows)
            if parsed is None:
                continue
            caption = None
            if buffer and len(buffer[-1].text.strip()) <= TITLE_MAX_CHARS:
                caption = buffer.pop().text.strip()      # short paragraph right above = title
            elif elements and isinstance(elements[-1], Document) and elements[-1].metadata.get("is_heading"):
                caption = elements[-1].page_content
            flush()
            has_header = parsed["has_header"] or (flagged and len(parsed["columns"]) > 1)
            columns, rows = parsed["columns"], parsed["rows"]
            header_row = parsed["header_row"]
            if flagged and not parsed["has_header"] and len(parsed["columns"]) > 1 and rows:
                header_row = rows[0]
                columns, rows = make_unique_columns(rows[0]), rows[1:]
            table = RawTable(
                rows=rows, columns=columns, has_header=has_header,
                title=parsed["title"], heading=caption, section=section,
                header_row=header_row, paged=False,
            )
            previous = elements[-1] if elements else None
            if (isinstance(previous, RawTable) and has_header and previous.has_header
                    and parsed["title"] is None and caption is None
                    and [norm_name(c) for c in previous.columns] == [norm_name(c) for c in columns]):
                previous.rows.extend(rows)               # table split by a page break
                previous.merged_parts += 1
            else:
                elements.append(table)

    flush()

    # embedded pictures (scans, charts, screenshots)
    try:
        from PIL import Image
        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                image = Image.open(io.BytesIO(rel.target_part.blob)).convert("RGB")
            except Exception:
                continue
            if image.width * image.height < 40_000:       # skip icons / bullets
                continue
            text = ocr_image(image)
            if len(text) >= 20:
                elements.append(_text_doc(text, file_path, "Embedded image",
                                          element_type="ocr_text", figure=True))
    except ImportError:
        logger.info("Pillow not installed - skipping embedded image OCR")

    documents, table_count = [], 0
    for element in elements:
        if isinstance(element, RawTable):
            table_count += 1
            documents.append(table_to_document(element, file_path, "docx", f"t{table_count:03d}"))
        else:
            documents.append(element)
    return documents
